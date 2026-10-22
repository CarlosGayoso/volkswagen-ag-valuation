import re

from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook

# =====================================================
# PROJECT PATHS
# =====================================================

PROJECT_ROOT = Path(__file__).resolve().parent.parent

EXCEL_FILE = (
    PROJECT_ROOT
    / "Finance_Evaluation_VW_v2.xlsx"
)


# =====================================================
# MANUAL MARKET DATA
# =====================================================
# This value is not currently stored in the workbook.
# Update it manually when the market price changes.

CURRENT_PRICE = 67.00


# =====================================================
# VALIDATION HELPERS
# =====================================================

def _numbers(worksheet, references):
    """Read and validate numeric values from Excel cells."""

    values = []

    for reference in references:
        cell_value = worksheet[reference].value
        cell_name = f"{worksheet.title}!{reference}"

        if cell_value is None:
            raise ValueError(
                f"Cell {cell_name} has no calculated value. "
                "Open the workbook in Excel, recalculate it, "
                "save it and rerun Streamlit."
            )

        if isinstance(cell_value, bool):
            raise TypeError(
                f"Expected a numeric value in {cell_name}, "
                f"but found {cell_value!r}."
            )

        if not isinstance(cell_value, (int, float)):
            raise TypeError(
                f"Expected a numeric value in {cell_name}, "
                f"but found {cell_value!r}."
            )

        values.append(float(cell_value))

    return values

def _years(worksheet, references):
    """Read years from numeric cells or text cells."""

    years = []

    for reference in references:
        value = worksheet[reference].value

        if isinstance(value, (int, float)):
            year = int(value)

        elif isinstance(value, str):
            text = value.replace("\n", " ")

            numeric_tokens = [
                token
                for token in text.split()
                if token.isdigit()
            ]

            if not numeric_tokens:
                raise ValueError(
                    f"Could not read a year from "
                    f"{worksheet.title}!{reference}: "
                    f"{value!r}"
                )

            year = int(numeric_tokens[-1])

        else:
            raise TypeError(
                f"Expected a year in "
                f"{worksheet.title}!{reference}, "
                f"but found {value!r}."
            )

        years.append(year)

    return years


# =====================================================
# EXCEL LOADER
# =====================================================

def _load_model(
    file_path: Path,
):
    """
    Load the Volkswagen valuation model using the
    last calculated values saved by Excel.
    """

    if not file_path.exists():
        raise FileNotFoundError(
            f"Workbook not found:\n{file_path}\n\n"
            "Place Finance_Evaluation_VW_v2.xlsx "
            "in the App_VW root folder."
        )

    workbook = load_workbook(
        filename=file_path,
        data_only=True,
        read_only=True,
    )

    required_sheets = {
        "SOTP Valuation",
        "DCF Scenarios",
        "DCF by Business",
    }

    missing_sheets = (
        required_sheets
        .difference(workbook.sheetnames)
    )

    if missing_sheets:

        workbook.close()

        raise KeyError(
            "Missing required Excel sheets: "
            + ", ".join(
                sorted(missing_sheets)
            )
        )

    sotp_worksheet = workbook[
        "SOTP Valuation"
    ]

    dcf_worksheet = workbook[
        "DCF Scenarios"
    ]

    dcf_business_worksheet = workbook[
        "DCF by Business"
    ]


    # =================================================
    # AUTOMOTIVE
    # =================================================

    automotive_years = _years(
        sotp_worksheet,
        [
            "C2",
            "D2",
            "E2",
            "F2",
            "G2",
            "H2",
        ],
    )

    automotive_df = pd.DataFrame(
        {
            "Year": automotive_years,

            "Revenue": _numbers(
                sotp_worksheet,
                [
                    "C4",
                    "D4",
                    "E4",
                    "F4",
                    "G4",
                    "H4",
                ],
            ),

            "EBIT": _numbers(
                sotp_worksheet,
                [
                    "C8",
                    "D8",
                    "E8",
                    "F8",
                    "G8",
                    "H8",
                ],
            ),

            "FCF": _numbers(
                sotp_worksheet,
                [
                    "C15",
                    "D15",
                    "E15",
                    "F15",
                    "G15",
                    "H15",
                ],
            ),
        }
    )


    # =================================================
    # CHINA JOINT VENTURES
    # =================================================

    china_years = _years(
        sotp_worksheet,
        [
            "C18",
            "D18",
            "E18",
            "F18",
            "G18",
            "H18",
        ],
    )

    china_df = pd.DataFrame(
        {
            "Year": china_years,

            "Net Income": _numbers(
                sotp_worksheet,
                [
                    "C20",
                    "D20",
                    "E20",
                    "F20",
                    "G20",
                    "H20",
                ],
            ),

            "P/E": _numbers(
                sotp_worksheet,
                [
                    "C21",
                    "D21",
                    "E21",
                    "F21",
                    "G21",
                    "H21",
                ],
            ),

            "Value": _numbers(
                sotp_worksheet,
                [
                    "C22",
                    "D22",
                    "E22",
                    "F22",
                    "G22",
                    "H22",
                ],
            ),
        }
    )


    # =================================================
    # FINANCIAL SERVICES
    # =================================================

    finance_years = _years(
        sotp_worksheet,
        [
            "C25",
            "D25",
            "E25",
            "F25",
            "G25",
            "H25",
        ],
    )

    finance_df = pd.DataFrame(
        {
            "Year": finance_years,

            "Book Value": _numbers(
                sotp_worksheet,
                [
                    "C27",
                    "D27",
                    "E27",
                    "F27",
                    "G27",
                    "H27",
                ],
            ),

            "P/B": _numbers(
                sotp_worksheet,
                [
                    "C28",
                    "D28",
                    "E28",
                    "F28",
                    "G28",
                    "H28",
                ],
            ),

            "Value": _numbers(
                sotp_worksheet,
                [
                    "C29",
                    "D29",
                    "E29",
                    "F29",
                    "G29",
                    "H29",
                ],
            ),
        }
    )


    # =================================================
    # 2026 SOTP
    # =================================================

    sotp_df = pd.DataFrame(
        {
            "Business": [
                "Automotive",
                "China JV",
                "Financial Services",
            ],

            "Value": _numbers(
                sotp_worksheet,
                [
                    "S3",
                    "S4",
                    "S5",
                ],
            ),
        }
    )


    # =================================================
    # DCF SCENARIOS
    # =================================================

    dcf_years = [
        int(value)
        for value in _numbers(
            dcf_worksheet,
            [
                "G3",
                "H3",
                "I3",
                "J3",
                "K3",
            ],
        )
    ]

    dcf_df = pd.DataFrame(
        {
            "Year": dcf_years,

            # DCF Scenarios:
            # Bear scenario stock price
            "Bear": _numbers(
                dcf_worksheet,
                [
                    "G12",
                    "H12",
                    "I12",
                    "J12",
                    "K12",
                ],
            ),

            # DCF Scenarios:
            # Base scenario stock price
            "Base": _numbers(
                dcf_worksheet,
                [
                    "G25",
                    "H25",
                    "I25",
                    "J25",
                    "K25",
                ],
            ),

            # DCF by Business:
            # Conservative stock price
            "Conservative": _numbers(
                dcf_business_worksheet,
                [
                    "G13",
                    "H13",
                    "I13",
                    "J13",
                    "K13",
                ],
            ),
        }
    )


    # =================================================
    # SUMMARY VALUES
    # =================================================

    summary_values = {
        "SOTP_VALUE": float(
            sotp_df["Value"].sum()
        ),

        "DCF_BASE": float(
            dcf_df.loc[
                dcf_df["Year"] == 2026,
                "Base",
            ].iloc[0]
        ),

        "DCF_CONSERVATIVE": float(
            dcf_df.loc[
                dcf_df["Year"] == 2026,
                "Conservative",
            ].iloc[0]
        ),
    }

    workbook.close()

    return (
        summary_values,
        sotp_df,
        automotive_df,
        china_df,
        finance_df,
        dcf_df,
    )


# =====================================================
# LOAD MODEL
# =====================================================

(
    _summary_values,
    sotp_df,
    automotive_df,
    china_df,
    finance_df,
    dcf_df,
) = _load_model(EXCEL_FILE)


# =====================================================
# EXPORTED KPI VALUES
# =====================================================

SOTP_VALUE = _summary_values[
    "SOTP_VALUE"
]

DCF_BASE = _summary_values[
    "DCF_BASE"
]

DCF_CONSERVATIVE = _summary_values[
    "DCF_CONSERVATIVE"
]


# =====================================================
# IMPLIED UPSIDE
# =====================================================

UPSIDE = (
    (DCF_BASE - CURRENT_PRICE)
    / CURRENT_PRICE
) * 100

CONSERVATIVE_UPSIDE = (
    (
        DCF_CONSERVATIVE
        - CURRENT_PRICE
    )
    / CURRENT_PRICE
) * 100


# =====================================================
# MODEL METADATA
# =====================================================

MODEL_LAST_MODIFIED = (
    EXCEL_FILE
    .stat()
    .st_mtime
)